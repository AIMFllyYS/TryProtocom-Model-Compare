"""导出 benchmark 常用表格与图表：CSV / Markdown / XLSX / PNG / 单页 HTML 报告。"""
from __future__ import annotations

import base64
import csv
import html
import io
import math
from collections import defaultdict
from pathlib import Path

from .common import dimensions

CJK_FONTS = ["Noto Sans CJK SC", "Noto Sans CJK JP", "Noto Sans CJK TC", "Source Han Sans SC", "PingFang SC", "Microsoft YaHei", "SimHei", "WenQuanYi Micro Hei", "Arial Unicode MS", "DejaVu Sans"]


def _fmt(v):
    if v is None:
        return ""
    if isinstance(v, float):
        if math.isinf(v):
            return "∞"
        return f"{v:.2f}".rstrip("0").rstrip(".") if abs(v) < 1000 else f"{v:.0f}"
    if isinstance(v, bool):
        return "是" if v else "否"
    return str(v)


def _csv(path: Path, rows: list[dict], fields: list[str] | None = None) -> None:
    fields = fields or (list(rows[0].keys()) if rows else [])
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: _fmt(r.get(k)) for k in fields})


def _md(rows: list[dict], fields: list[str], headers: list[str] | None = None) -> str:
    headers = headers or fields
    out = ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    for r in rows:
        out.append("| " + " | ".join(_fmt(r.get(k)) for k in fields) + " |")
    return "\n".join(out)


def tables(agg: dict) -> dict[str, tuple[list[dict], list[str], list[str]]]:
    """返回 {表名: (行, 字段, 中文表头)}，所有导出共用。"""
    dims = dimensions()
    did = [d["id"] for d in dims]
    dname = {d["id"]: d["name"] for d in dims}
    T = {}
    board = []
    for b in agg["board"]:
        r = {"rank": b["rank"], "entrant": b["entrant"], "quality": b["quality"], "ci": f"{_fmt(b['ci_low'])}–{_fmt(b['ci_high'])}",
             **{f"d_{k}": b["dims"].get(k) for k in did}, "pass_rate": b["pass_rate"], "suite_exp_cost_usd": b["suite_exp_cost_usd"],
             "suite_exp_min": b["suite_exp_min"], "runs": b["runs"], "complete": b["complete"], "na_ratio": b["na_ratio"]}
        if "composite" in b:
            r["composite"] = b["composite"]
        board.append(r)
    f = ["rank", "entrant", "quality", "ci", *[f"d_{k}" for k in did], "pass_rate", "suite_exp_cost_usd", "suite_exp_min", "runs", "complete", "na_ratio"]
    h = ["名次", "参赛者（模型 @ harness）", "质量总分", "95% 置信区间", *[dname[k] for k in did], "达标率", "整套期望成本($)", "整套期望用时(分)", "运行次数", "评分完整", "N/A 占比"]
    if agg["board"] and "composite" in agg["board"][0]:
        f.insert(3, "composite")
        h.insert(3, "综合分")
    T["leaderboard"] = (board, f, h)
    T["radar"] = ([{"entrant": b["entrant"], **{k: b["dims"].get(k) for k in did}} for b in agg["board"]], ["entrant", *did], ["参赛者", *[dname[k] for k in did]])

    tkeys = sorted({t["task"] for t in agg["tasks"]})
    mat = defaultdict(dict)
    for t in agg["tasks"]:
        mat[t["entrant"]][t["task"]] = f"{_fmt(t['mean'])} ± {_fmt(t['sd'])}"
    T["task_matrix"] = ([{"entrant": e, **v} for e, v in mat.items()], ["entrant", *tkeys], ["参赛者", *tkeys])

    T["tasks"] = (agg["tasks"], ["entrant", "task", "runs", "mean", "sd", "min", "max", "pass_rate", "pass_at_k", "gate_fail_runs", "complete", "low_confidence"],
                  ["参赛者", "题目", "运行次数", "均分", "标准差", "最低", "最高", "达标率", "pass@k", "门槛失败次数", "评分完整", "低可信"])
    T["efficiency"] = (agg["tasks"], ["entrant", "task", "mean_cost_usd", "mean_min", "pass_rate", "exp_cost_usd", "exp_min", "cost_score", "speed_score"],
                       ["参赛者", "题目", "单次平均费用($)", "单次平均用时(分)", "达标率", "期望成功成本($)", "期望达标用时(分)", "成本效率分", "速度分"])
    tiers = [{"entrant": t["entrant"], "task": t["task"], **{k: (round(v, 1) if v is not None else None) for k, v in t["tiers"].items()}} for t in agg["tasks"]]
    T["tiers"] = (tiers, ["entrant", "task", "basic", "advanced", "excellent", "clean"], ["参赛者", "题目", "基础项得分率", "进阶项得分率", "卓越项得分率", "扣分项（越高越干净）"])
    dimrows = []
    for t in agg["tasks"]:
        dimrows.append({"entrant": t["entrant"], "task": t["task"], **{k: t["dims"].get(k) for k in did}})
    T["task_dims"] = (dimrows, ["entrant", "task", *did], ["参赛者", "题目", *[dname[k] for k in did]])
    T["runs"] = (agg["runs"], list(agg["runs"][0].keys()) if agg["runs"] else [], list(agg["runs"][0].keys()) if agg["runs"] else [])
    T["item_scores_long"] = (agg["long"], ["run_id", "entrant", "task", "id", "dim", "tier", "method", "weight", "status", "s", "value"],
                             ["运行", "参赛者", "题目", "检查项", "维度", "层级", "方式", "权重", "状态", "得分(0-1)", "原始值"])
    T["item_analysis"] = (agg["items"], ["task", "item_id", "dim", "tier", "method", "desc", "mean", "between_sd", "within_sd", "n", "flag"],
                          ["题目", "检查项", "维度", "层级", "方式", "描述", "平均得分", "参赛者间标准差", "同一参赛者内标准差", "样本数", "标记"])
    T["skill_uplift"] = (agg["uplift"], ["entrant", "task", "A_mean", "C_mean", "uplift", "A_anim", "C_anim", "anim_uplift", "A_cost", "C_cost", "A_min", "C_min"],
                         ["参赛者", "题目", "A 组均分", "C 组均分", "增益", "A 组动画分", "C 组动画分", "动画增益", "A 组费用", "C 组费用", "A 组用时", "C 组用时"])
    T["failures"] = (agg["failures"], ["entrant", "task", "run_id", "total", "tags"], ["参赛者", "题目", "运行", "总分", "失败归因"])
    qd = [d for d in dims if d["kind"] == "quality"]
    T["coverage"] = (agg["coverage"], ["task", "name", *[d["id"] for d in qd], "items"], ["题目", "名称", *[d["name"] for d in qd], "检查项数"])
    return T


def _setup_mpl():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    have = {f.name for f in font_manager.fontManager.ttflist}
    chosen = [f for f in CJK_FONTS if f in have and f != "DejaVu Sans"]
    if not chosen:  # 让 fontconfig 找一个中文字体文件
        import shutil
        import subprocess
        if shutil.which("fc-match"):
            fp = subprocess.run(["fc-match", "-f", "%{file}", ":lang=zh"], capture_output=True, text=True).stdout.strip()
            if fp:
                try:
                    font_manager.fontManager.addfont(fp)
                    chosen = [font_manager.FontProperties(fname=fp).get_name()]
                except Exception:  # noqa: BLE001
                    pass
    plt.rcParams["font.sans-serif"] = chosen + ["DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    return plt


PALETTE = ["#2f5bd3", "#d9622b", "#1f9e74", "#b8398f", "#7a6a12", "#4a8fb3", "#c23b3b", "#6b4fc2"]


def radar_png(agg: dict, path: Path, only: str | None = None) -> None:
    import numpy as np
    plt = _setup_mpl()
    dims = dimensions()
    labels = [d["name"] for d in dims]
    ang = np.linspace(0, 2 * np.pi, len(labels), endpoint=False).tolist()
    ang += ang[:1]
    fig = plt.figure(figsize=(7, 7), dpi=150)
    ax = fig.add_subplot(111, polar=True)
    for i, b in enumerate(agg["board"]):
        if only and b["entrant"] != only:
            continue
        vals = [(b["dims"].get(d["id"]) or 0) for d in dims]
        vals += vals[:1]
        c = PALETTE[i % len(PALETTE)]
        ax.plot(ang, vals, color=c, linewidth=2, label=b["entrant"])
        ax.fill(ang, vals, color=c, alpha=0.10)
    ax.set_xticks(ang[:-1])
    ax.set_xticklabels(labels, fontsize=11)
    ax.set_ylim(0, 100)
    ax.set_yticks([20, 40, 60, 80, 100])
    ax.set_yticklabels(["20", "40", "60", "80", "100"], fontsize=8, color="#666")
    ax.grid(color="#ccc", linewidth=0.6)
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.06), ncol=2, fontsize=9, frameon=False)
    missing = [f"{b['entrant']}：{'、'.join(d['name'] for d in dims if b['dims'].get(d['id']) is None)}" for b in agg["board"]
               if (not only or b["entrant"] == only) and any(b["dims"].get(d["id"]) is None for d in dims)]
    if missing:
        fig.text(0.5, -0.02, "无数据的维度按 0 绘制 —— " + "；".join(missing), ha="center", fontsize=8, color="#777", wrap=True)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def pareto_png(agg: dict, path: Path) -> list[str]:
    plt = _setup_mpl()
    pts = [(b["entrant"], b["suite_exp_cost_usd"], b["quality"]) for b in agg["board"] if b["suite_exp_cost_usd"] not in (None, math.inf) and b["quality"] is not None]
    front = []
    for e, c, q in pts:
        if not any((c2 <= c and q2 >= q) and (c2 < c or q2 > q) for _, c2, q2 in pts):
            front.append(e)
    fig, ax = plt.subplots(figsize=(8, 5.5), dpi=150)
    for i, (e, c, q) in enumerate(pts):
        on = e in front
        ax.scatter([c], [q], s=90 if on else 50, color=PALETTE[i % len(PALETTE)], edgecolor="black" if on else "none", zorder=3)
        ax.annotate(e, (c, q), textcoords="offset points", xytext=(6, 6), fontsize=8.5)
    fp = sorted([(c, q) for e, c, q in pts if e in front])
    if len(fp) > 1:
        ax.plot([p[0] for p in fp], [p[1] for p in fp], color="#888", linestyle="--", linewidth=1, zorder=2)
    if pts:
        ax.set_xscale("log")
    ax.set_xlabel("完成整套题的期望成本（美元，对数刻度）")
    ax.set_ylabel("质量总分")
    ax.set_title("质量 × 成本（黑边为帕累托前沿）", fontsize=11)
    ax.grid(color="#e3e3e3", linewidth=0.6)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    return front


def task_bars_png(agg: dict, path: Path) -> None:
    import numpy as np
    plt = _setup_mpl()
    tkeys = sorted({t["task"] for t in agg["tasks"]})
    ents = [b["entrant"] for b in agg["board"]]
    fig, ax = plt.subplots(figsize=(max(8, len(tkeys) * 1.2), 5), dpi=150)
    w = 0.8 / max(len(ents), 1)
    for i, e in enumerate(ents):
        ms = [next((t["mean"] for t in agg["tasks"] if t["entrant"] == e and t["task"] == tk), 0) for tk in tkeys]
        sds = [next((t["sd"] for t in agg["tasks"] if t["entrant"] == e and t["task"] == tk), 0) for tk in tkeys]
        x = np.arange(len(tkeys)) + i * w - 0.4 + w / 2
        ax.bar(x, ms, w, yerr=sds, color=PALETTE[i % len(PALETTE)], label=e, capsize=2)
    ax.set_xticks(np.arange(len(tkeys)))
    ax.set_xticklabels(tkeys)
    ax.set_ylim(0, 100)
    ax.set_ylabel("题目得分（均值 ± 标准差）")
    ax.legend(fontsize=8, frameon=False, ncol=2)
    ax.grid(axis="y", color="#e3e3e3", linewidth=0.6)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def export_all(agg: dict, out: Path, formats: set[str]) -> list[Path]:
    out.mkdir(parents=True, exist_ok=True)
    T = tables(agg)
    written = []
    if "csv" in formats:
        for name, (rows, f, h) in T.items():
            p = out / f"{name}.csv"
            _csv(p, rows, f)
            written.append(p)
    if "md" in formats:
        p = out / "leaderboard.md"
        parts = [f"# {agg.get('benchmark')} {agg.get('version')} 排行榜\n"]
        for name, title in (("leaderboard", "总榜"), ("radar", "雷达维度"), ("task_matrix", "题目得分矩阵（均值 ± 标准差）"),
                            ("efficiency", "效率"), ("skill_uplift", "Skill 增益（A/C 对照）"), ("failures", "失败归因")):
            rows, f, h = T[name]
            parts.append(f"## {title}\n\n{_md(rows, f, h)}\n")
        p.write_text("\n".join(parts), encoding="utf-8")
        written.append(p)
    charts = {}
    if "png" in formats and agg["board"]:
        radar_png(agg, out / "radar.png")
        for i, b in enumerate(agg["board"]):
            radar_png(agg, out / f"radar_{i + 1:02d}.png", only=b["entrant"])
        pareto_png(agg, out / "pareto.png")
        task_bars_png(agg, out / "task_scores.png")
        charts = {"radar": out / "radar.png", "pareto": out / "pareto.png", "tasks": out / "task_scores.png"}
        written += list(charts.values()) + sorted(out.glob("radar_*.png"))
    if "xlsx" in formats:
        p = out / "benchmark.xlsx"
        _xlsx(T, p, charts)
        written.append(p)
    if "html" in formats:
        p = out / "report.html"
        _report_html(agg, T, charts, p)
        written.append(p)
    return written


SHEET_TITLES = {"leaderboard": "总榜", "radar": "雷达维度", "task_matrix": "题目矩阵", "tasks": "题目统计", "efficiency": "效率",
                "tiers": "分层得分", "task_dims": "题目×维度", "stability": "稳定性", "skill_uplift": "Skill增益", "failures": "失败归因",
                "item_analysis": "题目分析", "coverage": "覆盖矩阵", "runs": "运行明细", "item_scores_long": "检查项明细"}


def _xlsx(T, path: Path, charts: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="E8EDF7")
    for name, (rows, f, h) in T.items():
        ws = wb.create_sheet(SHEET_TITLES.get(name, name)[:31])
        ws.append(h)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for r in rows:
            vals = []
            for k in f:
                v = r.get(k)
                vals.append("∞" if isinstance(v, float) and math.isinf(v) else v)
            ws.append(vals)
        ws.freeze_panes = "B2"
        for i, k in enumerate(f, 1):
            width = max([len(str(h[i - 1]))] + [len(_fmt(r.get(k))) for r in rows[:200]]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(max(width, 8), 48)
    if charts:
        ws = wb.create_sheet("图表")
        row = 1
        for key in ("radar", "pareto", "tasks"):
            if key in charts and Path(charts[key]).exists():
                img = XLImage(str(charts[key]))
                img.width, img.height = img.width * 0.55, img.height * 0.55
                ws.add_image(img, f"A{row}")
                row += int(img.height / 20) + 2
    wb.save(path)


def _report_html(agg, T, charts, path: Path) -> None:
    def img(p):
        if not p or not Path(p).exists():
            return ""
        return f'<img alt="" src="data:image/png;base64,{base64.b64encode(Path(p).read_bytes()).decode()}">'

    def table(name):
        rows, f, h = T[name]
        head = "".join(f"<th>{html.escape(x)}</th>" for x in h)
        body = "".join("<tr>" + "".join(f"<td>{html.escape(_fmt(r.get(k)))}</td>" for k in f) + "</tr>" for r in rows)
        return f'<div class="tw"><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>'

    sections = [("总榜", "leaderboard"), ("雷达维度", "radar"), ("题目得分矩阵", "task_matrix"), ("分层得分（区分度来源）", "tiers"),
                ("效率", "efficiency"), ("Skill 增益", "skill_uplift"), ("失败归因", "failures"), ("题目分析（用于迭代题库）", "item_analysis")]
    body = "".join(f"<section><h2>{t}</h2>{table(n)}</section>" for t, n in sections)
    doc = f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(str(agg.get('benchmark')))} 报告</title><style>
:root{{--bg:#f7f8fa;--fg:#1c2230;--muted:#5e6778;--line:#dfe3ea;--card:#fff}}
@media (prefers-color-scheme:dark){{:root{{--bg:#12151b;--fg:#e5e8ee;--muted:#9aa3b2;--line:#2c333e;--card:#1a1f27;color-scheme:dark}}}}
body{{margin:0;background:var(--bg);color:var(--fg);font:14px/1.55 system-ui,"PingFang SC","Microsoft YaHei",sans-serif}}
main{{max-width:1180px;margin:0 auto;padding-block:24px;padding-inline:16px;display:grid;gap:20px}}
h1{{font-size:22px;margin:0}} h2{{font-size:16px;margin:0 0 10px}}
section{{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:16px;min-width:0}}
.charts{{display:grid;grid-template-columns:repeat(auto-fit,minmax(320px,1fr));gap:16px}} .charts img{{width:100%;background:#fff;border-radius:6px}}
.tw{{overflow-x:auto}} table{{border-collapse:collapse;font-size:12.5px;font-variant-numeric:tabular-nums;min-width:100%}}
th,td{{border-bottom:1px solid var(--line);padding:5px 8px;text-align:left;white-space:nowrap}} th{{color:var(--muted);font-weight:600}}
.muted{{color:var(--muted)}}</style></head><body><main>
<header><h1>{html.escape(str(agg.get('benchmark')))} {html.escape(str(agg.get('version')))}</h1>
<p class="muted">冻结日期：{html.escape(str(agg.get('frozen_date') or '未设置'))} · 参赛者 {len(agg['board'])} 个 · 运行 {len(agg['runs'])} 次。置信区间重叠的名次记为并列（“=”）。</p></header>
<section><h2>图表</h2><div class="charts">{img(charts.get('radar'))}{img(charts.get('pareto'))}{img(charts.get('tasks'))}</div></section>
{body}</main></body></html>"""
    path.write_text(doc, encoding="utf-8")
